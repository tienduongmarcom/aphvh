"""Tiền và vay (lượt 0.14, cờ `kh_aph_tien_vay`, mặc định tắt). Đặc tả docs/dac-ta-0.14.md mục 3–7.

- Tài khoản ngân hàng khai trong khach.json (mục `ngan_hang`: đơn vị, ngân hàng, 4 số cuối, loại tiền): tài khoản chi tiết
  dưới nhóm của 1121 / 1122 + thẻ Bank Account. Chỉ thêm cái thiếu.
- Báo cáo "So quy va tien gui" (S07 / S08): tồn sau từng phiếu.
- Thu / chi hộ qua tài khoản cá nhân giám đốc: bút toán nháp 1388 / 3388 đối tượng Nhân viên; báo cáo "Cong no voi giam doc".
- Khế ước vay (DocType "APH Khe uoc vay"), ô khế ước trên dòng bút toán (bắt buộc với 341x), báo cáo "Du no vay theo khe uoc",
  trích lãi tháng (nháp Nợ 6351 / Có 3352), dòng tiền 7/30 ngày cộng gốc vay đến hạn.
- Đối chiếu sao kê Excel / CSV với sổ (mã giao dịch, số tiền, ngày ±2): khớp / lệch / trùng / chưa ghi sổ. Không tự ghi sổ.
Quy tắc thuần ở lab9x.tien_vay.
"""
import hashlib
import json

import frappe
from frappe.utils import add_days, cint, flt, get_first_day, get_last_day, getdate, nowdate
from lab9x import tien_vay as T

FLAG = "kh_aph_tien_vay"
DT_VAY = "APH Khe uoc vay"
F_KU = "kh_aph_khe_uoc"
F_LAI = "kh_aph_lai_thang"
F_HO = "kh_aph_thu_chi_ho"
F_MAU = "kh_aph_mau_sao_ke"
WORKSPACE = "Tiền và vay"
VAI_TRO_KT = {"Accounts User", "Accounts Manager", "System Manager"}
LOAI_HO = {
    # loại -> (bên Nợ, bên Có); "nv1388"/"nv3388": tài khoản của giám đốc; "kh": phải thu khách; "ncc": phải trả NCC;
    # "doi_ung": tài khoản chi phí / tiền kế toán chọn
    "Giám đốc thu hộ tiền khách": ("nv1388", "kh"),
    "Giám đốc chi hộ trả nhà cung cấp": ("ncc", "nv3388"),
    "Giám đốc chi hộ chi phí": ("doi_ung", "nv3388"),
    "Giám đốc nộp lại tiền thu hộ": ("doi_ung", "nv1388"),
    "Công ty hoàn tiền chi hộ": ("nv3388", "doi_ung"),
}
TK_LAI = {"6351": ("Chi phí lãi vay", "635", "Expense"), "3352": ("Chi phí lãi vay phải trả", "335", "Liability")}

FIELDS = {
    "Journal Entry Account": [
        {"fieldname": F_KU, "label": "Khế ước vay", "fieldtype": "Link", "options": DT_VAY, "insert_after": "party",
         "description": "Bắt buộc với tài khoản vay 341x; gắn cả dòng lãi 635 / 335 để theo dõi lãi theo khế ước."}],
    "Journal Entry": [
        {"fieldname": F_LAI, "label": "Trích lãi vay tháng", "fieldtype": "Data", "insert_after": "voucher_type",
         "read_only": 1, "no_copy": 1, "hidden": 1},
        {"fieldname": F_HO, "label": "Thu / chi hộ qua tài khoản cá nhân", "fieldtype": "Data", "insert_after": F_LAI,
         "read_only": 1, "no_copy": 1, "depends_on": F_HO}],
    "Bank": [
        {"fieldname": F_MAU, "label": "Mẫu cột sao kê (JSON)", "fieldtype": "Code", "options": "JSON",
         "insert_after": "bank_name",
         "description": "Ví dụ {\"dong_tieu_de\": 1, \"ngay\": \"Ngày GD\", \"ma\": \"Số tham chiếu\", \"noi_dung\": \"Diễn giải\", "
                        "\"ghi_no\": \"Ghi nợ\", \"ghi_co\": \"Ghi có\"}. Cột ghi theo tên tiêu đề hoặc số thứ tự (bắt đầu 1)."}],
}


def enabled():
    return cint(frappe.conf.get(FLAG))


def _can_bat():
    if not enabled():
        frappe.throw("Tiền và vay 0.14 chưa được bật")


def _vai_tro(ktt=False):
    from kh_aph.constants import ROLE_NO_COST
    roles = set(frappe.get_roles())
    can = {"Accounts Manager", "System Manager"} if ktt else VAI_TRO_KT
    if frappe.session.user != "Administrator" and (not (roles & can) or ROLE_NO_COST in roles):
        frappe.throw("Cần quyền " + ("kế toán trưởng" if ktt else "kế toán"), frappe.PermissionError)


def schema():
    """Gọi từ uat_candidate_setup.schema (UAT và bộ cài production). Không gắn after_install/after_migrate."""
    if not enabled():
        frappe.throw("Cần bật cờ kh_aph_tien_vay trước khi cài")
    frappe.reload_doc("kh_aph", "doctype", frappe.scrub(DT_VAY))
    from frappe.custom.doctype.custom_field.custom_field import create_custom_fields
    create_custom_fields(FIELDS)
    tk = [tai_khoan_lai(c) for c in frappe.get_all("Company", pluck="name")]
    _workspace()
    return {"fields": sum(len(v) for v in FIELDS.values()), "tai_khoan_lai": [x for x in tk if x],
            "workspace": WORKSPACE}


def _workspace():
    loi_vao = [
        ("Tiền thực nhận", "DocType", "APH Funds Receipt", {"docstatus": 1}),
        ("Tiền chưa rõ khách", "DocType", "APH Funds Receipt", {"docstatus": 1, "customer": ["is", "not set"]}),
        ("Phân bổ tiền và cọc còn nháp", "DocType", "APH Funds Action", {"docstatus": 0}),
        ("Tuyến thu tiền đã duyệt", "DocType", "APH Funds Route", {"docstatus": 1}),
        ("Bàn giao tiền nhà máy chờ nhận", "DocType", "APH Phieu quy nha may", {"docstatus": 1, "loai": "Bàn giao tiền", "trang_thai": "Chờ nhận"}),
        ("Bàn giao tiền có chênh lệch", "DocType", "APH Phieu quy nha may", {"docstatus": 1, "trang_thai": "Chênh lệch chờ điều tra"}),
        ("Sổ quỹ và tiền gửi", "Report", "So quy va tien gui", None),
        ("Khế ước vay", "DocType", DT_VAY, {"trang_thai": "Đang vay"}),
        ("Dư nợ vay theo khế ước", "Report", "Du no vay theo khe uoc", None),
        ("Công nợ với giám đốc", "Report", "Cong no voi giam doc", None),
        ("Tài khoản ngân hàng", "DocType", "Bank Account", None),
        ("Bút toán nháp", "DocType", "Journal Entry", {"docstatus": 0}),
        ("Dòng tiền 7/30 ngày", "Report", "Dong tien 7 30 ngay", None),
    ]
    if cint(frappe.conf.get("kh_aph_cash_custody")):
        loi_vao.append(("Rà người giữ tiền mặt", "Report", "Ra nguoi giu tien mat", None))
    loi_vao = [x for x in loi_vao if frappe.db.exists(x[1], x[2])]
    if cint(frappe.conf.get("kh_aph_cash_closing")):
        loi_vao += [("Xác minh quỹ", "DocType", "APH Xac minh quy", {}),
                    ("Chốt ca tiền", "DocType", "APH Chot ca tien", {}),
                    ("Ca tiền chênh lệch", "DocType", "APH Chot ca tien", {"docstatus": 1, "status": "Có chênh lệch"})]
    noi_dung = [{"id": "tv-tieu-de", "type": "header", "data": {"text": f"<span class=\"h4\"><b>{WORKSPACE}</b></span>", "col": 12}}]
    noi_dung += [{"id": f"tv-{i}", "type": "shortcut", "data": {"shortcut_name": x[0], "col": 4}} for i, x in enumerate(loi_vao)]
    ws = frappe.get_doc("Workspace", WORKSPACE) if frappe.db.exists("Workspace", WORKSPACE) else frappe.new_doc("Workspace")
    ws.update({"label": WORKSPACE, "title": WORKSPACE, "public": 1, "module": "Kh Aph", "icon": "money-coins-1",
               "content": json.dumps(noi_dung, ensure_ascii=False)})
    ws.set("shortcuts", [{"label": nhan, "type": loai, "link_to": dich, "color": "Grey",
                          **({"stats_filter": json.dumps(loc), "format": "{} " + nhan.lower()} if loc else {})}
                         for nhan, loai, dich, loc in loi_vao])
    ws.set("roles", [{"role": r} for r in sorted(VAI_TRO_KT) if frappe.db.exists("Role", r)])
    ws.flags.ignore_permissions = True
    ws.save() if not ws.is_new() else ws.insert()


# ---------------------------------------------------------------- tài khoản

def tk_so(company, so_hieu):
    """Tài khoản chi tiết đúng số hiệu (hoặc chi tiết duy nhất bắt đầu bằng số hiệu)."""
    v = frappe.db.get_value("Account", {"company": company, "account_number": so_hieu, "is_group": 0}, "name")
    if v:
        return v
    ds = frappe.get_all("Account", {"company": company, "account_number": ["like", so_hieu + "%"], "is_group": 0}, pluck="name")
    return ds[0] if len(ds) == 1 else None


def tai_khoan_lai(company):
    """6351 (dưới 635) và 3352 (dưới 335) nếu đơn vị có nhóm 635 / 335 (APH, VH). Chỉ thêm, chạy lại an toàn."""
    tao = []
    for so_hieu, (ten, nhom, rt) in TK_LAI.items():
        if frappe.db.exists("Account", {"company": company, "account_number": so_hieu}):
            continue
        cha = frappe.db.get_value("Account", {"company": company, "account_number": nhom, "is_group": 1}, "name")
        if not cha:
            continue
        frappe.get_doc({"doctype": "Account", "account_name": ten, "account_number": so_hieu, "company": company,
                        "parent_account": cha, "root_type": rt, "is_group": 0}).insert(ignore_permissions=True)
        tao.append(so_hieu)
    return {"company": company, "tao": tao} if tao else None


def _doc_cfg(duong_dan=None, cfg=None):
    if cfg is None:
        from kh_aph.khach_hang import doc_cau_hinh
        cfg = doc_cau_hinh(duong_dan or frappe.get_app_path("kh_aph", "..", "khach", "aph-vh", "khach.json"))
    return cfg


def dung_ngan_hang(duong_dan=None, cfg=None):
    """Dựng tài khoản ngân hàng theo mục `ngan_hang` của khach.json. Chạy tay (Administrator / kế toán trưởng):

        bench --site <site> execute kh_aph.tien_vay.dung_ngan_hang --kwargs '{"duong_dan": "<khach.json>"}'

    Số tài khoản đầy đủ không có trong cấu hình: kế toán nhập vào ô "Bank Account No" trên thẻ Bank Account."""
    _can_bat()
    _vai_tro(ktt=True)
    cfg = _doc_cfg(duong_dan, cfg)
    abbrs = {c["abbr"] for c in cfg.get("cong_ty") or []}
    loi = T.kiem_ngan_hang(cfg.get("ngan_hang") or [], abbrs)
    if loi:
        frappe.throw("Mục ngan_hang không hợp lệ: " + "; ".join(loi))
    kq = {"tao_tai_khoan": [], "tao_the": [], "da_co": [], "canh_bao": []}
    for r in cfg.get("ngan_hang") or []:
        cty = frappe.db.get_value("Company", {"abbr": r["don_vi"]}, "name")
        if not cty:
            kq["canh_bao"].append(f"Chưa có đơn vị {r['don_vi']} trên site")
            continue
        lt = r.get("loai_tien") or "VND"
        ten = T.ten_tk(r["ngan_hang"], r["so_cuoi"], lt)
        tk = frappe.db.get_value("Account", {"company": cty, "account_name": ten}, "name")
        if tk:
            kq["da_co"].append(tk)
        else:
            goc = T.LOAI_TIEN[lt]
            moc = frappe.db.get_value("Account", {"company": cty, "account_number": goc}, ["name", "is_group", "parent_account"],
                                      as_dict=True) or frappe.db.get_value(
                "Account", {"company": cty, "account_number": "1121"}, ["name", "is_group", "parent_account"], as_dict=True)
            if not moc:
                kq["canh_bao"].append(f"{cty}: không thấy tài khoản {goc} / 1121")
                continue
            cha = moc.name if moc.is_group else moc.parent_account
            so_co = set(frappe.get_all("Account", {"company": cty, "account_number": ["like", goc + "%"]}, pluck="account_number"))
            so_moi = T.so_hieu_moi(goc, so_co)
            d = frappe.get_doc({"doctype": "Account", "account_name": ten, "account_number": so_moi, "company": cty,
                                "parent_account": cha, "root_type": "Asset", "account_type": "Bank", "is_group": 0,
                                "account_currency": lt})
            d.insert(ignore_permissions=True)
            tk = d.name
            kq["tao_tai_khoan"].append(tk)
        if not frappe.db.exists("Bank", r["ngan_hang"].strip()):
            frappe.get_doc({"doctype": "Bank", "bank_name": r["ngan_hang"].strip()}).insert(ignore_permissions=True)
        if not frappe.db.exists("Bank Account", {"account": tk}):
            # tên thẻ Bank Account là duy nhất toàn site (account_name - ngân hàng): thêm viết tắt đơn vị
            # (APH và VH cùng có MB đuôi 9999)
            ba = frappe.get_doc({"doctype": "Bank Account", "account_name": f"{ten} {r['don_vi']}", "bank": r["ngan_hang"].strip(),
                                 "account": tk,
                                 "company": cty, "is_company_account": 1, "mask": r["so_cuoi"]})
            ba.insert(ignore_permissions=True)
            kq["tao_the"].append(ba.name)
    frappe.db.commit()
    return kq


# ---------------------------------------------------------------- sổ quỹ / sổ tiền gửi

def so_quy(company, account, tu_ngay, den_ngay):
    frappe.get_doc("Company", company).check_permission("read")
    frappe.get_doc("Account", account).check_permission("read")
    tu, den = getdate(tu_ngay), getdate(den_ngay)
    if tu > den:
        frappe.throw("Từ ngày phải trước đến ngày")
    a = frappe.db.get_value("Account", account, ["company", "account_type", "is_group", "account_currency"], as_dict=True)
    if not a or a.company != company or a.account_type not in ("Cash", "Bank"):
        frappe.throw("Chọn tài khoản tiền mặt / ngân hàng của đúng đơn vị")
    dau = flt(frappe.db.sql("""SELECT SUM(debit_in_account_currency - credit_in_account_currency) FROM `tabGL Entry`
        WHERE account=%s AND is_cancelled=0 AND posting_date<%s""", (account, tu))[0][0])
    gl = frappe.db.sql("""SELECT posting_date AS ngay, voucher_type, voucher_no AS so, remarks AS dien_giai,
            party, SUM(debit_in_account_currency) AS thu, SUM(credit_in_account_currency) AS chi
        FROM `tabGL Entry` WHERE account=%s AND is_cancelled=0 AND posting_date BETWEEN %s AND %s
        GROUP BY posting_date, voucher_type, voucher_no ORDER BY posting_date, creation, voucher_no""",
                       (account, tu, den), as_dict=True)
    for r in gl:
        r.loai = "Phiếu thu" if flt(r.thu) >= flt(r.chi) else "Phiếu chi"
        net = flt(r.thu) - flt(r.chi)
        r.thu, r.chi = (net, 0) if net >= 0 else (0, -net)
    dong, thu, chi, cuoi = T.so_ton(dau, [dict(r) for r in gl])
    return {"dau_ky": dau, "dong": dong, "thu": thu, "chi": chi, "cuoi_ky": cuoi, "tien": a.account_currency}


# ---------------------------------------------------------------- thu / chi hộ của giám đốc

@frappe.whitelist(methods=["POST"])
def thu_chi_ho(company, loai, nhan_vien, so_tien, noi_dung, khach=None, nha_cung_cap=None, tk_doi_ung=None, ngay=None):
    """Lập bút toán NHÁP; kế toán kiểm rồi xác nhận. Không ghi vào 112 của công ty (tiền ở tài khoản cá nhân)."""
    _can_bat()
    _vai_tro()
    if loai not in LOAI_HO:
        frappe.throw("Loại thu / chi hộ không hợp lệ")
    tien = flt(so_tien, 2)
    if tien <= 0 or not (noi_dung or "").strip():
        frappe.throw("Cần số tiền dương và nội dung")
    if not frappe.db.exists("Employee", nhan_vien):
        frappe.throw("Chọn người thu / chi hộ trong danh mục Nhân viên")
    from erpnext.accounts.party import get_party_account

    def dong(ben):
        if ben in ("nv1388", "nv3388"):
            tk = tk_so(company, ben[2:])
            if not tk:
                frappe.throw(f"{company} chưa có tài khoản chi tiết {ben[2:]}")
            return {"account": tk, "party_type": "Employee", "party": nhan_vien}
        if ben == "kh":
            if not khach:
                frappe.throw("Chọn khách")
            return {"account": get_party_account("Customer", khach, company), "party_type": "Customer", "party": khach}
        if ben == "ncc":
            if not nha_cung_cap:
                frappe.throw("Chọn nhà cung cấp")
            return {"account": get_party_account("Supplier", nha_cung_cap, company), "party_type": "Supplier",
                    "party": nha_cung_cap}
        a = frappe.db.get_value("Account", tk_doi_ung, ["company", "is_group"], as_dict=True) if tk_doi_ung else None
        if not a or a.company != company or a.is_group:
            frappe.throw("Chọn tài khoản đối ứng chi tiết của đúng đơn vị")
        return {"account": tk_doi_ung}

    no, co = LOAI_HO[loai]
    je = frappe.get_doc({"doctype": "Journal Entry", "voucher_type": "Journal Entry", "company": company,
                         "posting_date": str(getdate(ngay or nowdate())), F_HO: loai,
                         "user_remark": f"{loai}: {noi_dung.strip()}",
                         "accounts": [{**dong(no), "debit_in_account_currency": tien},
                                      {**dong(co), "credit_in_account_currency": tien}]})
    je.insert()
    return {"journal_entry": je.name}


def cong_no_giam_doc(company=None, nhan_vien=None, den_ngay=None):
    allowed = tuple(frappe.get_list("Company", pluck="name", limit_page_length=0))
    if company and company not in allowed:
        frappe.throw("Không có quyền xem đơn vị đã chọn", frappe.PermissionError)
    if not allowed:
        return []
    dk, gt = " AND g.company IN %(companies)s", {"d": str(getdate(den_ngay or nowdate())),
        "companies": (company,) if company else allowed}
    if nhan_vien:
        dk += " AND g.party=%(nv)s"
        gt["nv"] = nhan_vien
    rows = frappe.db.sql(f"""SELECT g.company, g.party AS nhan_vien,
            SUM(CASE WHEN a.account_number LIKE '1388%%' THEN g.debit - g.credit ELSE 0 END) AS phai_thu,
            SUM(CASE WHEN a.account_number LIKE '3388%%' THEN g.credit - g.debit ELSE 0 END) AS phai_tra
        FROM `tabGL Entry` g JOIN `tabAccount` a ON a.name=g.account
        WHERE g.is_cancelled=0 AND g.party_type='Employee' AND g.posting_date<=%(d)s
          AND (a.account_number LIKE '1388%%' OR a.account_number LIKE '3388%%') {dk}
        GROUP BY g.company, g.party ORDER BY g.party, g.company""", gt, as_dict=True)
    for r in rows:
        r.ten = frappe.db.get_value("Employee", r.nhan_vien, "employee_name")
        r.chenh = flt(r.phai_thu) - flt(r.phai_tra)
    return rows


# ---------------------------------------------------------------- khế ước vay

def validate_but_toan(doc, method=None):
    """Journal Entry validate: dòng tài khoản vay 341x phải gắn khế ước đúng đơn vị và đúng tài khoản của khế ước."""
    if not enabled() or not frappe.get_meta("Journal Entry Account").has_field(F_KU):
        return
    for d in doc.get("accounts") or []:
        so = str(frappe.get_cached_value("Account", d.account, "account_number") or "")
        ku = d.get(F_KU)
        if so.startswith("341") and not ku:
            frappe.throw(f"Dòng {d.idx}: tài khoản vay {d.account} phải chọn khế ước vay", title="Thiếu khế ước")
        if ku:
            k = frappe.db.get_value(DT_VAY, ku, ["company", "tai_khoan"], as_dict=True)
            if not k or k.company != doc.company:
                frappe.throw(f"Dòng {d.idx}: khế ước {ku} không thuộc đơn vị {doc.company}")
            if so.startswith("341") and d.account != k.tai_khoan:
                frappe.throw(f"Dòng {d.idx}: khế ước {ku} ghi trên tài khoản {k.tai_khoan}")


def _dong_khe_uoc(ten, den=None, tien_to=("341",)):
    dk = " OR ".join("a.account_number LIKE %s" for _ in tien_to)
    gt = [ten] + [t + "%" for t in tien_to]
    them = ""
    if den:
        them = " AND je.posting_date<=%s"
        gt.append(str(getdate(den)))
    return frappe.db.sql(f"""SELECT je.posting_date AS ngay, SUM(d.credit - d.debit) AS tien
        FROM `tabJournal Entry Account` d JOIN `tabJournal Entry` je ON je.name=d.parent JOIN `tabAccount` a ON a.name=d.account
        WHERE je.docstatus=1 AND d.{F_KU}=%s AND ({dk}){them} GROUP BY je.posting_date ORDER BY je.posting_date""",
                         gt, as_dict=True)


def bien_dong(ten, den=None):
    """[(ngày, gốc nhận + / trả −)] của khế ước, từ bút toán đã xác nhận."""
    return [(r.ngay, flt(r.tien)) for r in _dong_khe_uoc(ten, den)]


def du_no_vay(company=None, ngay=None, chi_dang_vay=True):
    hom_nay = getdate(ngay or nowdate())
    if company:
        frappe.get_doc("Company", company).check_permission("read")
    loc = {"company": company} if company else {}
    if chi_dang_vay:
        loc["trang_thai"] = "Đang vay"
    out = []
    for k in frappe.get_list(DT_VAY, loc, ["name", "company", "so_khe_uoc", "ben_cho_vay", "so_tien", "lai_suat", "ngay_vay",
                                           "ngay_den_han", "ky_han", "trang_thai"], order_by="ngay_den_han", limit_page_length=0):
        bd = bien_dong(k.name, hom_nay)
        nhan = sum(t for _n, t in bd if t > 0)
        tra = -sum(t for _n, t in bd if t < 0)
        lai_cp = -sum(flt(r.tien) for r in _dong_khe_uoc(k.name, hom_nay, ("635",)))
        lai_phai_tra = sum(flt(r.tien) for r in _dong_khe_uoc(k.name, hom_nay, ("335",)))
        out.append({**k, "goc_nhan": nhan, "da_tra": tra, "du_no": round(nhan - tra, 2), "lai_chi_phi": lai_cp,
                    "lai_phai_tra": lai_phai_tra, "nhom_han": T.nhom_han(k.ngay_den_han, hom_nay)})
    return out


@frappe.whitelist(methods=["POST"])
def trich_lai(company, den_ngay=None):
    """Bút toán NHÁP trích lãi tháng chứa den_ngay: Nợ 6351 / Có 3352 theo từng khế ước (dư nợ thực từng ngày × lãi suất),
    trừ lãi đã ghi thẳng vào 635 trong tháng. Mỗi đơn vị mỗi tháng một bút toán."""
    _can_bat()
    _vai_tro(ktt=True)
    den = getdate(den_ngay or nowdate())
    tu, cuoi = get_first_day(den), get_last_day(den)
    if den != cuoi and den < getdate(nowdate()):
        den = cuoi
    khoa = f"{company}|{den:%Y-%m}"
    da = frappe.db.get_value("Journal Entry", {F_LAI: khoa, "docstatus": ["<", 2]})
    if da:
        frappe.throw(f"Tháng {den:%m/%Y} đã có bút toán trích lãi {da}")
    no, co = tk_so(company, "6351"), tk_so(company, "3352")
    if not (no and co):
        frappe.throw(f"{company} chưa có tài khoản 6351 / 3352 (chạy bộ cài tiền và vay)")
    dong = []
    for k in frappe.get_all(DT_VAY, {"company": company, "trang_thai": "Đang vay"}, ["name", "lai_suat", "so_khe_uoc"]):
        lai, _du = T.lai_ky(bien_dong(k.name), flt(k.lai_suat), tu, den)
        da_ghi = flt(frappe.db.sql(f"""SELECT SUM(d.debit - d.credit) FROM `tabJournal Entry Account` d
            JOIN `tabJournal Entry` je ON je.name=d.parent JOIN `tabAccount` a ON a.name=d.account
            WHERE je.docstatus=1 AND d.{F_KU}=%s AND a.account_number LIKE '635%%' AND je.posting_date BETWEEN %s AND %s""",
                                   (k.name, tu, den))[0][0])
        tien = round(lai - da_ghi)
        if tien > 0:
            dong += [{"account": no, "debit_in_account_currency": tien, F_KU: k.name},
                     {"account": co, "credit_in_account_currency": tien, F_KU: k.name}]
    if not dong:
        return {"journal_entry": None, "ly_do": "Không có lãi cần trích"}
    je = frappe.get_doc({"doctype": "Journal Entry", "voucher_type": "Journal Entry", "company": company,
                         "posting_date": str(den), F_LAI: khoa, "user_remark": f"Trích lãi vay tháng {den:%m/%Y}",
                         "accounts": dong})
    je.insert()
    return {"journal_entry": je.name, "so_khe_uoc": len(dong) // 2}


def khoan_vay_den_han(company, ngay):
    """Cho dòng tiền 7/30 ngày (quan_tri.dong_tien_cong_ty): [(ngày đến hạn, gốc còn nợ + lãi chưa trả), chi tiết]."""
    if not enabled() or not frappe.db.exists("DocType", DT_VAY):
        return []
    out = []
    for r in du_no_vay(company, ngay):
        con = flt(r["du_no"]) + flt(r["lai_phai_tra"])
        if con > 0:
            out.append((r["ngay_den_han"], con, r))
    return out


# ---------------------------------------------------------------- mã giao dịch / sao kê

def ma_da_ghi(tk, ma, bo_qua=None, lock=False):
    """Chứng từ đã ghi mã giao dịch `ma` trên tài khoản ngân hàng `tk` (phiếu thu / chi, bút toán, phiếu thực thu 0.10)."""
    ma = (ma or "").strip()
    if not ma:
        return None
    suffix = " FOR UPDATE" if lock else ""
    v = frappe.db.sql("""SELECT name FROM `tabPayment Entry` WHERE docstatus<2 AND reference_no=%s
        AND (paid_to=%s OR paid_from=%s) AND name!=%s LIMIT 1""" + suffix, (ma, tk, tk, bo_qua or ""))
    if v:
        return v[0][0]
    v = frappe.db.sql("""SELECT je.name FROM `tabJournal Entry` je JOIN `tabJournal Entry Account` d ON d.parent=je.name
        WHERE je.docstatus<2 AND je.cheque_no=%s AND d.account=%s AND je.name!=%s LIMIT 1""" + suffix, (ma, tk, bo_qua or ""))
    if v:
        return v[0][0]
    if frappe.db.exists("DocType", "APH Funds Receipt"):
        key = hashlib.sha256((tk + "\0" + ma).encode()).hexdigest()
        v = frappe.db.get_value("APH Funds Receipt", {"external_key": key, "docstatus": ["<", 2], "name": ["!=", bo_qua or ""]}, for_update=lock)
        if v:
            return v
    if frappe.db.exists("DocType", "APH Funds Action"):
        v = frappe.db.sql("""SELECT name FROM `tabAPH Funds Action` WHERE docstatus<2 AND name!=%s
            AND ((bank_account=%s AND external_reference=%s) OR (owner_bank_account=%s AND owner_external_reference=%s))
            AND action_type IN ('Hoàn tiền', 'Trả tiền thu hộ') LIMIT 1""" + suffix,
            (bo_qua or "", tk, ma, tk, ma))
        if v:
            return v[0][0]
    return None


def _bang_tu_tep(file_url):
    f = frappe.get_doc("File", {"file_url": file_url})
    f.check_permission("read")
    noi_dung = f.get_content()
    ten = (f.file_name or file_url).lower()
    if ten.endswith((".xlsx", ".xlsm")):
        import io
        from openpyxl import load_workbook
        # không dùng read_only: tệp MB ghi sai kích thước trang tính, chế độ read_only chỉ đọc được 1 dòng
        wb = load_workbook(io.BytesIO(noi_dung), data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    if ten.endswith(".xls"):
        import xlrd  # có sẵn trong bench Frappe
        sh = xlrd.open_workbook(file_contents=noi_dung).sheet_by_index(0)
        return [sh.row_values(i) for i in range(sh.nrows)]
    if ten.endswith(".csv"):
        import csv
        import io
        txt = noi_dung.decode("utf-8-sig") if isinstance(noi_dung, bytes) else noi_dung
        mau = txt[:2000]
        dau = ";" if mau.count(";") > mau.count(",") else ","
        return [r for r in csv.reader(io.StringIO(txt), delimiter=dau)]
    frappe.throw("Chỉ nhận tệp .xlsx, .xls hoặc .csv")


def _so_sach(tk, tu, den):
    """Chứng từ đã ghi trên tài khoản ngân hàng trong khoảng ngày: mã giao dịch từ phiếu thu / chi, bút toán, phiếu thực thu."""
    gl = frappe.db.sql("""SELECT voucher_type, voucher_no, posting_date, SUM(debit_in_account_currency - credit_in_account_currency) AS tien
        FROM `tabGL Entry` WHERE account=%s AND is_cancelled=0 AND posting_date BETWEEN %s AND %s
        GROUP BY voucher_type, voucher_no, posting_date""", (tk, tu, den), as_dict=True)
    out = []
    for g in gl:
        ma = ""
        if g.voucher_type == "Payment Entry":
            ma = frappe.db.get_value("Payment Entry", g.voucher_no, "reference_no") or ""
        elif g.voucher_type == "Journal Entry":
            ma = frappe.db.get_value("Journal Entry", g.voucher_no, "cheque_no") or ""
            if not ma and frappe.db.exists("DocType", "APH Funds Receipt"):
                ma = frappe.db.get_value("APH Funds Receipt", {"journal_entry": g.voucher_no}, "external_reference") or ""
        out.append({"chung_tu": g.voucher_no, "loai": g.voucher_type, "ma": ma.strip(), "ngay": g.posting_date,
                    "so_tien": flt(g.tien, 2)})
    return out


@frappe.whitelist(methods=["POST"])
def doi_chieu_sao_ke(bank_account, file_url, anh_xa=None, luu_mau=0):
    """Đọc tệp sao kê của một tài khoản ngân hàng (thẻ Bank Account) và so với sổ. Chỉ đọc; không ghi sổ."""
    _can_bat()
    _vai_tro()
    ba = frappe.get_doc("Bank Account", bank_account)
    ba.check_permission("read")
    if not ba.account or not ba.company:
        frappe.throw("Thẻ tài khoản ngân hàng chưa gắn tài khoản kế toán / đơn vị")
    account_doc = frappe.get_doc("Account", ba.account)
    account_doc.check_permission("read")
    if account_doc.company != ba.company:
        frappe.throw("Tài khoản kế toán không thuộc đơn vị của thẻ ngân hàng")
    bank = None
    if cint(luu_mau) and anh_xa:
        bank = frappe.get_doc("Bank", ba.bank)
        bank.check_permission("write")
    if anh_xa:
        mau = frappe.parse_json(anh_xa) if isinstance(anh_xa, str) else anh_xa
    else:
        mau = frappe.parse_json(frappe.db.get_value("Bank", ba.bank, F_MAU) or "{}") or T.mau_ngan_hang(ba.bank)
    if not mau:
        frappe.throw(f"Ngân hàng {ba.bank} chưa có mẫu cột sao kê: khai mẫu trên thẻ Ngân hàng hoặc trong hộp thoại")
    bang = _bang_tu_tep(file_url)
    dong, loi = T.doc_sao_ke(bang, mau)
    co_san = T.mau_ngan_hang(ba.bank)
    if not dong and not anh_xa and co_san and co_san != mau:  # mẫu đã lưu không hợp tệp: thử mẫu có sẵn của ngân hàng
        dong, loi = T.doc_sao_ke(bang, co_san)
    if loi and not dong:
        frappe.throw("; ".join(loi))
    if cint(luu_mau) and anh_xa:
        bank.set(F_MAU, json.dumps(mau, ensure_ascii=False))
        bank.save()
    if not dong:
        return {"dong": [], "loi": loi, "tong": {}}
    tu, den = min(d["ngay"] for d in dong), max(d["ngay"] for d in dong)
    ket = T.khop_sao_ke(dong, _so_sach(ba.account, add_days(tu, -3), add_days(den, 3)))
    tong = {}
    for r in ket:
        tong[r["trang_thai"]] = tong.get(r["trang_thai"], 0) + 1
        r["ngay"] = str(r["ngay"])
    return {"dong": ket, "loi": loi, "tong": tong, "tai_khoan": ba.account, "company": ba.company}


def has_permission(doc, ptype=None, user=None):
    """Frappe 16: chỉ True / False. Khế ước vay: cờ tắt thì chặn; hồ sơ không xem giá vốn (Dũng) không xem."""
    if not enabled():
        return False
    from kh_aph.constants import ROLE_NO_COST
    return ROLE_NO_COST not in frappe.get_roles(user)
