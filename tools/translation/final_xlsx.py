import json, collections, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
H='/home/claude/aphvh/tools/translation/'
o=json.load(open(H+'final_rows.json'))
ns={}; src=open(H+'build_proposals.py').read()
exec(src.split("# ---- Thuật ngữ chuẩn")[1].split("PROPOSE =")[0].split('\n',1)[1], ns)
KEYS=list(ns['APPROVED'])+[k for k in ns['GLOSS'] if k not in ns['APPROVED']]
by_en={r['src']:r for r in o}
F='Arial'; hdr=Font(name=F,bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='1F3864'); thin=Side(style='thin',color='BFBFBF')
red=PatternFill('solid',fgColor='FDE9E7')
wb=openpyxl.Workbook()
def sheet(title, cols, widths, rows, hl=None):
    ws=wb.create_sheet(title); ws.append(cols)
    for c in ws[1]: c.font=hdr; c.fill=fill; c.alignment=Alignment(wrap_text=True,vertical='center')
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=Font(name=F,size=10); c.alignment=Alignment(wrap_text=True,vertical='top'); c.border=Border(bottom=thin)
        if hl and hl(row): 
            for c in row: c.fill=red
    for i,w in enumerate(widths,1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.freeze_panes='C2'; ws.auto_filter.ref=ws.dimensions
ch=[r for r in o if r['changed2']]
cnt=collections.Counter((r['log'] or ['Chuẩn hóa thuật ngữ / chữ hoa'])[0] for r in ch)
biz=[r for r in o if r['log'] and r['log'][0]=='Sai nghiệp vụ']
ws=wb.active; ws.title='Đọc trước'
L=[('Bản dịch chốt – màn hình kế toán trưởng (ERPNext APH)',),(),('Ngày','29/09/2026'),
('Phạm vi','1.067 chuỗi hiển thị trên màn hình Dung: mã hàng, nhà cung cấp, đơn mua, phiếu nhập mua, phiếu kho, điều chỉnh theo kiểm kê, phân bổ chi phí nhập hàng, kho, công ty/khóa sổ, tài sản, bút toán, nút – trạng thái – báo cáo – menu.'),
('Cách làm','Vòng 1: đối chiếu bản đang hiện (ERPNext 16.36.1 chính thức) với thuật ngữ đã chốt, ERPNext2VI 2.9 và kho ERPNext2VI. Vòng 2: 6 người rà song song đọc từng chuỗi theo đúng chứng từ và trường của nó (sai nghiệp vụ, khó hiểu, nhất quán). Vòng 3: một người kiểm định độc lập đọc lại toàn bộ bản chốt, soát nhất quán giữa các màn hình và nhãn trùng chữ.'),
('Kết quả',f'{len(ch)} chuỗi đổi so với bản đang hiện, trong đó {len(biz)} lỗi sai nghiệp vụ (sheet 1). Không còn chuỗi nào dịch Stock thành cổ phiếu, Rate thành tỷ giá (trừ tỷ giá thật), Billing thành thanh toán, Credit thành ghi nợ; không mất placeholder {{0}}.'),
('Thuật ngữ đã chốt','Xác nhận / Đã xác nhận · Nháp · Hủy · Sửa đổi · Dòng · Đối tượng · Đơn giá vốn · Cộng tiền hàng · Tổng tiền thanh toán · Định mức nguyên vật liệu · Nguyên vật liệu · Ký hiệu chứng từ · Điều chỉnh theo kiểm kê · Nhập số dư đầu kỳ · Phiếu nhập mua hàng · Mặt hàng · Thêm dòng (xem sheet 3).'),
('Còn giới hạn','Mỗi chuỗi tiếng Anh chỉ có một bản dịch cho toàn hệ thống. Ví dụ "Opening Stock" hiện "Nhập số dư đầu kỳ" cả trên phiếu điều chỉnh lẫn ô số dư đầu kỳ khi tạo mã hàng; "Stock" để "Kho" vì dùng chung cho tên phân hệ.'),
(),('Loại thay đổi','Số chuỗi')]+sorted(cnt.items(),key=lambda x:-x[1])+[(),('Tiếp theo','Đóng gói vào app aphvh 0.2.2 → thử 8080 → 8081 → cloud (mỗi người Ctrl+Shift+R). Sau đó mở rộng màn hình Dũng, rồi bán hàng, tiền, công nợ theo cùng bảng thuật ngữ.')]
for l in L: ws.append(list(l))
ws['A1'].font=Font(name=F,bold=True,size=14)
for r in ws.iter_rows(min_row=2):
    for c in r: c.font=Font(name=F,size=10,bold=(c.column==1)); c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=24; ws.column_dimensions['B'].width=110
sheet('1. Lỗi sai nghiệp vụ đã sửa',['STT','Tiếng Anh','Đang hiện','Bản chốt','Vì sao sai','Màn hình'],[5,34,34,36,60,28],
      [[i+1,r['src'],r['cur'] or '(tiếng Anh)',r['final'],r['log'][1],'; '.join(r['screens'])] for i,r in enumerate(sorted(biz,key=lambda r:r['prio']))])
sheet('2. Toàn bộ bản chốt',['STT','Tiếng Anh','Đang hiện','Bản chốt','Đổi?','Loại','Lý do','Màn hình','Ưu tiên'],[5,34,32,36,7,16,50,28,8],
      [[i+1,r['src'],r['cur'] or '(tiếng Anh)',r['final'],'Có' if r['changed2'] else '',(r['log'] or ['Chuẩn hóa' if r['changed2'] else ''])[0],(r['log'] or ['',''])[1],'; '.join(r['screens']),r['prio']] for i,r in enumerate(sorted(o,key=lambda r:(r['prio'],r['screens'][0] if r['screens'] else '',r['src'])))],
      hl=lambda row: row[5].value=='Sai nghiệp vụ')
g=[[i+1,k,((lambda v: 'Tài khoản '+v[3:] if v.startswith('TK ') else v)(by_en[k]['final'] if k in by_en else ns['GLOSS'].get(k, ns['APPROVED'].get(k)))),('Đã chốt' if k in ns['APPROVED'] else 'Chốt sau rà soát')] for i,k in enumerate(KEYS)]
sheet('3. Bảng thuật ngữ',['STT','Tiếng Anh','Tiếng Việt chốt','Trạng thái'],[5,40,44,18],g)
out='/mnt/user-data/outputs/Ban_dich_chot_man_hinh_Dung_APH.xlsx'; wb.save(out); print(out,len(ch),len(biz),cnt)
